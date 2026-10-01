"""Verify saved CN workbooks, including explicit shells with no neighbors."""

from pathlib import Path
import tempfile
import unittest

from ase import Atoms
from openpyxl import load_workbook

from neighbor_core import ClusterNeighbor


class CoordinationExportTests(unittest.TestCase):
    def test_workbook_retains_empty_shell_ids_and_metadata(self):
        cluster = ClusterNeighbor()
        cluster.load_xyz(from_file=False, atom_object=Atoms(
            "CuO2", positions=[[0, 0, 0], [1, 0, 0], [2, 0, 0]]))
        cluster.get_CN_summary_all(bond_range=4, printit=False,
                                   shell_edges={"Cu-O": [0, 1, 2, 3, 4]})
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "coordination.xlsx"
            cluster.write_to_excel(path)
            workbook = load_workbook(path, data_only=True)
            try:
                values = list(workbook["Cu-O"].values)
                rows = [dict(zip(values[0], row)) for row in values[1:]]
                self.assertEqual([row["Shell"] for row in rows], [1, 2, 3, 4])
                self.assertEqual([row["Average CN"] for row in rows], [0, 1, 1, 0])
                self.assertEqual([row["Bond length [A]"] for row in rows], [None, "1.000", "2.000", None])
                self.assertEqual([row["Lower bound [A]"] for row in rows], [0, 1, 2, 3])
                self.assertEqual([row["Upper bound [A]"] for row in rows], [1, 2, 3, 4])
                self.assertEqual([row["Shell method"] for row in rows], ["edges"] * 4)
                self.assertEqual([row["Distance std [A]"] for row in rows], [None, 0, 0, None])
            finally:
                workbook.close()

    def test_empty_cluster_exports_a_readable_workbook(self):
        cluster = ClusterNeighbor()
        cluster.load_xyz(from_file=False, atom_object=Atoms())
        cluster.get_CN_summary_all(printit=False)
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "empty.xlsx"
            cluster.write_to_excel(path)
            workbook = load_workbook(path, data_only=True)
            try:
                self.assertEqual(workbook.sheetnames, ["Summary"])
                self.assertEqual(list(workbook.active.values), [("Bond length [A]", "Average CN")])
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
